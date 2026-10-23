# -*- coding: utf-8 -*-
"""把 .docx 原卷拆成 正文.txt（含 [TABLE] 标记）+ media/ 图片目录。

用法:
    python extract_docx.py "<原卷.docx>" "<输出目录>"

输出:
    <输出目录>/正文.txt     按文档顺序，段落原样，表格用 [TABLE i rows=x cols=y] 包住
    <输出目录>/media/       内嵌图片，保留原文件名（image1.png / image10.jpeg ...）

注意:
    - .doc（老二进制）先用 Word COM 转 docx:
        $w = New-Object -ComObject Word.Application
        $d = $w.Documents.Open($src); $d.SaveAs2($dst, 12); $d.Close($false)
      用 SaveAs2，不要用 SaveAs([ref]...)。
    - zip 里 word/media/ 含目录条目，basename 为空，必须过滤，否则 open() 会抛 FileNotFoundError。
"""
import os
import sys
import zipfile

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    src, outdir = sys.argv[1], sys.argv[2]
    os.makedirs(outdir, exist_ok=True)

    d = docx.Document(src)
    lines, ti = [], 0
    for child in d.element.body.iterchildren():
        if child.tag.endswith("}p"):
            t = Paragraph(child, d).text
            if t.strip():
                lines.append(t)
        elif child.tag.endswith("}tbl"):
            tb = Table(child, d)
            lines.append(f"[TABLE {ti} rows={len(tb.rows)} cols={len(tb.columns)}]")
            for r in tb.rows:
                lines.append(" | ".join(c.text.replace("\n", " / ").strip() for c in r.cells))
            lines.append("[/TABLE]")
            ti += 1

    txt_path = os.path.join(outdir, "正文.txt")
    with open(txt_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))

    media_dir = os.path.join(outdir, "media")
    os.makedirs(media_dir, exist_ok=True)
    with zipfile.ZipFile(src) as z:
        names = [n for n in z.namelist()
                 if n.startswith("word/media/") and not n.endswith("/")]
        for n in names:
            with open(os.path.join(media_dir, os.path.basename(n)), "wb") as fh:
                fh.write(z.read(n))

    print(f"段落 {len(d.paragraphs)} | 表格 {len(d.tables)} | 内嵌图 {len(d.inline_shapes)} | 媒体文件 {len(names)}")
    print(f"正文 -> {txt_path}")
    print(f"图片 -> {media_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
